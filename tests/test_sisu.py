import json
import logging

import pytest
from openpyxl import Workbook

from sauf_scraper.cli import main
from sauf_scraper.geral.sisu import ler_sisu, nota, relatar_codigos, turno

CABECALHO = [
    "EDICAO", "CO_IES", "NO_IES", "CO_IES_CURSO", "NO_CURSO", "DS_GRAU", "DS_TURNO",
    "NO_MUNICIPIO_CAMPUS", "TP_MOD_CONCORRENCIA", "TIPO_CONCORRENCIA",
    "QT_VAGAS_OFERTADAS", "NU_NOTACORTE", "QT_INSCRICAO",
]  # fmt: skip

# Linhas inventadas no layout da planilha do Sisu: espaços sobrando em NO_CURSO e
# TIPO_CONCORRENCIA, cotas (L/V) que devem ser ignoradas e corte 0 = sem nota.
LINHAS = [
    [2025, 57, "UNIV A", 5000540, "BIOMEDICINA   ", "Bacharelado", "Integral",
     "Maringá", "A", "AC     ", 2, 761.72, 610],
    [2025, 57, "UNIV A", 5000540, "BIOMEDICINA   ", "Bacharelado", "Integral",
     "Maringá", "L", "LB_PPI ", 1, 700.10, 80],
    [2025, 57, "UNIV A", 5000540, "BIOMEDICINA   ", "Bacharelado", "Integral",
     "Maringá", "V", "V", 1, 650.00, 12],
    [2025, 57, "UNIV A", 3421, "QUÍMICA", "Licenciatura", " NOTURNO ",
     "Maringá", " A ", "AC", 6, 0, 13],
    [2025, 9, "UNIV FORA", 1234, "DIREITO", "Bacharelado", "Matutino",
     "Londrina", "A", "AC", 10, 780.0, 900],
]  # fmt: skip


@pytest.fixture
def planilha(tmp_path):
    pasta = Workbook()
    pasta.active.title = "Dicionário de dados"
    aba = pasta.create_sheet("inscricao_2025_1")
    aba.append(CABECALHO)
    for linha in LINHAS:
        aba.append(linha)
    caminho = tmp_path / "sisu-2025.xlsx"
    pasta.save(caminho)
    return caminho


def test_so_ampla_concorrencia_das_ies_do_escopo(planilha):
    lote, relatorio = ler_sisu(planilha, [57])

    assert lote.ano == 2025
    assert [(x.codigo_emec_curso, x.turno, x.nota) for x in lote.linhas] == [
        (5000540, "Integral", 761.72),
        (3421, "Noturno", None),
    ]
    assert relatorio.encontradas == {57}


def test_campos_de_cada_linha(planilha):
    linha = ler_sisu(planilha, [57])[0].linhas[0]

    assert linha.model_dump(by_alias=True) == {
        "codigoEmecInstituicao": 57,
        "codigoEmecCurso": 5000540,
        "turno": "Integral",
        "vagas": 2,
        "nota": 761.72,
        "inscricoes": 610,
    }


@pytest.mark.parametrize(
    ("valor", "esperado"),
    [
        (" NOTURNO ", "Noturno"),
        ("integral", "Integral"),
        ("EaD", "EaD"),
        ("Vespertino", "Vespertino"),
    ],
)
def test_turno_normalizado(valor, esperado):
    assert turno(valor) == esperado


@pytest.mark.parametrize(
    ("valor", "esperado"), [(0, None), (0.0, None), (None, None), (712.4, 712.4)]
)
def test_corte_zero_vira_null(valor, esperado):
    assert nota(valor) == esperado


def test_aba_sem_prefixo_e_erro(tmp_path):
    pasta = Workbook()
    pasta.active.title = "outra"
    pasta.save(tmp_path / "x.xlsx")

    with pytest.raises(ValueError, match="inscricao_"):
        ler_sisu(tmp_path / "x.xlsx", [57])


def test_relatorio_compara_com_a_tabela_de_casamento(planilha, caplog):
    lote, _ = ler_sisu(planilha, [57])

    with caplog.at_level(logging.INFO):
        relatar_codigos(lote, {57: {5000540, 9999}})

    registro = next(r for r in caplog.records if "Sisu" in r.getMessage())
    assert registro.dados["noSisuForaDaTabela"] == [3421]
    assert registro.dados["naTabelaSemSisu"] == [9999]
    assert registro.dados["semNota"] == 1


def test_comando_sisu_dry_run(planilha, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr("sauf_scraper.cli.codigos_emec", lambda: [57])
    monkeypatch.setattr("sauf_scraper.cli.codigos_por_ies", lambda: {57: {5000540}})

    assert main(["sisu", "--arquivo", str(planilha), "--dry-run"]) == 0

    conteudo = json.loads((tmp_path / "output" / "sisu-2025.json").read_text(encoding="utf-8"))
    assert conteudo["ano"] == 2025
    assert len(conteudo["linhas"]) == 2

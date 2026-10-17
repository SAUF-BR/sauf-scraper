"""Base comum dos importadores gerais (D-48): planilha local → IES do escopo → destino.

Cada importador (igc.py, e depois o Sisu) só sabe ler as próprias colunas. Ler o xlsx,
filtrar as IES com spider, relatar o que não casou e mandar o resultado para o JSON ou
para a API fica aqui.
"""

import logging
from collections.abc import Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from sauf_scraper.client.destinos import ApiSauf, ArquivoJson
from sauf_scraper.core.config import Settings
from sauf_scraper.core.models import RespostaIngestao, SaufModel

logger = logging.getLogger(__name__)


def ler_planilha(arquivo: Path, aba: str) -> Iterator[dict[str, Any]]:
    """Linhas da aba como dicts. Os cabeçalhos das planilhas do INEP vêm com espaços nas
    pontas (" Código da IES"), por isso o strip."""
    pasta = load_workbook(arquivo, read_only=True, data_only=True)
    try:
        linhas = pasta[aba].iter_rows(values_only=True)
        cabecalho = [str(c).strip() if c is not None else "" for c in next(linhas)]
        for linha in linhas:
            if any(valor is not None for valor in linha):
                yield dict(zip(cabecalho, linha, strict=False))
    finally:
        pasta.close()


@dataclass
class Relatorio:
    """O que casou e o que não casou entre a planilha e as IES do escopo."""

    fonte: str
    escopo: list[int]
    encontradas: set[int] = field(default_factory=set)

    @property
    def ausentes(self) -> list[int]:
        """IES com spider que não aparecem na planilha."""
        return sorted(set(self.escopo) - self.encontradas)

    def registrar(self) -> None:
        nivel = logging.WARNING if self.ausentes else logging.INFO
        logger.log(
            nivel,
            "casamento por código e-MEC",
            extra={
                "dados": {
                    "fonte": self.fonte,
                    "escopo": self.escopo,
                    "encontradas": sorted(self.encontradas),
                    "semCorrespondencia": self.ausentes,
                }
            },
        )


def do_escopo(
    linhas: Iterator[dict[str, Any]], coluna_codigo: str, relatorio: Relatorio
) -> Iterator[dict[str, Any]]:
    """Só as linhas das IES com spider (P-35 c). Anota no relatório quais foram achadas."""
    escopo = set(relatorio.escopo)
    for linha in linhas:
        codigo = inteiro(linha.get(coluna_codigo))
        if codigo in escopo:
            relatorio.encontradas.add(codigo)
            yield linha


def inteiro(valor: Any) -> int | None:
    if isinstance(valor, int):
        return valor
    texto = str(valor).strip() if valor is not None else ""
    return int(texto) if texto.isdigit() else None


def entregar(
    lote: SaufModel, nome_arquivo: str, endpoint: str, settings: Settings, dry_run: bool
) -> RespostaIngestao | None:
    """--dry-run grava output/<nome_arquivo>; sem ele, envia para a API."""
    if dry_run:
        ArquivoJson(settings.output_dir).gravar(nome_arquivo, lote)
        return None
    return ApiSauf(settings).postar(endpoint, lote)

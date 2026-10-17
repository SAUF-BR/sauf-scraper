"""G3: notas de corte do Sisu, só ampla concorrência (D-45, D-50).

Planilha "Inscrições e notas de corte" do Sisu (gov.br/mec/pt-br/sisu/documentos), aba
"inscricao_<edição>". Uma linha por curso × turno × modalidade de concorrência; ficam só as
linhas de ampla concorrência (TP_MOD_CONCORRENCIA "A") das IES do escopo. NU_NOTACORTE 0
quer dizer que não houve nota de corte e vira null.
"""

import logging
from pathlib import Path

from openpyxl import load_workbook
from pydantic import Field

from sauf_scraper.core.models import SaufModel
from sauf_scraper.geral.base import Relatorio, do_escopo, inteiro, ler_planilha

logger = logging.getLogger(__name__)

ENDPOINT = "/api/v1/ingestao/sisu"
PREFIXO_ABA = "inscricao_"
AMPLA_CONCORRENCIA = "A"

# Mesma grafia usada no site e nas opções das ofertas ("Integral", "Noturno"...).
_TURNOS = {
    "integral": "Integral",
    "matutino": "Matutino",
    "vespertino": "Vespertino",
    "noturno": "Noturno",
    "ead": "EaD",
}


class LinhaSisu(SaufModel):
    codigo_emec_instituicao: int = Field(gt=0)
    codigo_emec_curso: int = Field(gt=0)
    turno: str = Field(min_length=1)
    vagas: int | None = Field(default=None, ge=0)
    nota: float | None = Field(default=None, gt=0, le=1000)
    inscricoes: int | None = Field(default=None, ge=0)


class LoteSisu(SaufModel):
    ano: int
    linhas: list[LinhaSisu]


def ler_sisu(arquivo: Path, escopo: list[int]) -> tuple[LoteSisu, Relatorio]:
    relatorio = Relatorio(fonte=arquivo.name, escopo=sorted(escopo))
    anos: set[int] = set()
    linhas = []
    for linha in do_escopo(ler_planilha(arquivo, aba_de_inscricao(arquivo)), "CO_IES", relatorio):
        if _texto(linha.get("TP_MOD_CONCORRENCIA")) != AMPLA_CONCORRENCIA:
            continue
        anos.add(int(linha["EDICAO"]))
        linhas.append(
            LinhaSisu(
                codigo_emec_instituicao=inteiro(linha["CO_IES"]),
                codigo_emec_curso=inteiro(linha["CO_IES_CURSO"]),
                turno=turno(linha["DS_TURNO"]),
                vagas=inteiro(linha.get("QT_VAGAS_OFERTADAS")),
                nota=nota(linha.get("NU_NOTACORTE")),
                inscricoes=inteiro(linha.get("QT_INSCRICAO")),
            )
        )
    if len(anos) > 1:
        raise ValueError(f"planilha com mais de uma edição: {sorted(anos)}")
    if not anos:
        raise ValueError("nenhuma linha de ampla concorrência das IES do escopo")
    return LoteSisu(ano=anos.pop(), linhas=linhas), relatorio


def aba_de_inscricao(arquivo: Path) -> str:
    """O nome da aba muda por edição (ex.: "inscricao_2025_1")."""
    pasta = load_workbook(arquivo, read_only=True)
    try:
        abas = [nome for nome in pasta.sheetnames if nome.startswith(PREFIXO_ABA)]
    finally:
        pasta.close()
    if len(abas) != 1:
        raise ValueError(f"esperava uma aba '{PREFIXO_ABA}...', achei {abas}")
    return abas[0]


def turno(valor) -> str:
    texto = _texto(valor)
    return _TURNOS.get(texto.lower(), texto.title())


def nota(valor) -> float | None:
    """0 (ou vazio) = sem nota de corte → None."""
    if valor is None or isinstance(valor, str) and not valor.strip():
        return None
    numero = float(str(valor).replace(",", "."))
    return numero if numero > 0 else None


def relatar_codigos(lote: LoteSisu, confirmados: dict[int, set[int]]) -> None:
    """Compara os códigos de curso do Sisu com os confirmados nas tabelas de casamento.

    `confirmados`: código e-MEC da IES → códigos de curso com confirmado = sim.
    """
    for ies, codigos_csv in sorted(confirmados.items()):
        linhas = [x for x in lote.linhas if x.codigo_emec_instituicao == ies]
        codigos_sisu = {x.codigo_emec_curso for x in linhas}
        logger.info(
            "códigos de curso do Sisu × tabela de casamento",
            extra={
                "dados": {
                    "ies": ies,
                    "ano": lote.ano,
                    "linhas": len(linhas),
                    "semNota": sum(1 for x in linhas if x.nota is None),
                    "noSisuForaDaTabela": sorted(codigos_sisu - codigos_csv),
                    "naTabelaSemSisu": sorted(codigos_csv - codigos_sisu),
                }
            },
        )


def _texto(valor) -> str:
    return str(valor).strip() if valor is not None else ""
